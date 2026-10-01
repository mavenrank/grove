"""Reviewed work, typed pack and verified media must survive every boundary."""
import hashlib
import io
import json
from copy import deepcopy

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app.config import settings
from app.db import db
from ingestion.cli import main
from ingestion.importing import install_media, prepare_import, write_candidate
from ingestion.pack_contract import content_digest, validate_pack
from ingestion.pipeline import Pipeline


def test_approved_pack_cannot_contain_a_draft_lesson(candidate):
    payload = deepcopy(candidate[2])
    source = payload["concepts"][0]["learning_segments"][0]["source"]
    citation = {k: source[k] for k in ("deck_id", "source_file", "source_path", "source_hash", "slide_id", "slide_number")}
    payload["concepts"][0]["lesson"] = {"schema_version": 1, "id": "pilot", "title": "Pilot",
        "introduction": "Review this lesson before publication.", "review_status": "draft", "assets": [],
        "sections": [{"id": "baseline", "title": "Baseline", "stage": "baseline", "blocks": [
            {"type": "text", "id": "intro", "paragraphs": ["Candidate teaching text"], "sources": [citation]}]}]}
    payload["content_sha256"] = content_digest(payload)
    with pytest.raises(ValueError, match="draft lessons"):
        validate_pack(payload)


@pytest.fixture
def candidate(tmp_path):
    source = tmp_path / "source"
    source.mkdir()
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(2)).text = (
        "A ratio a:b divides a whole into a plus b equal parts; one part equals total divided by the sum.")
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(2)).text = (
        "Question 1\nDivide twelve items in the ratio of two to one. What is the larger share?\nA) 8\nB) 4")
    slide.notes_slide.notes_text_frame.text = "Answer: A. Divide by three and multiply by two."
    prs.save(source / "RATIO.pptx")
    pipeline = Pipeline(source, tmp_path / "work")
    assert pipeline.run_all() == 0
    assert pipeline.run_pack("1.0", True) == 0
    path = pipeline.work_dir / "packs/grove-ingested-1.0.json"
    return pipeline, path, json.loads(path.read_text(encoding="utf-8"))


@pytest.mark.parametrize("failure", ["truthy_approval", "legacy_schema", "missing_review", "missing_summary",
                                    "wrong_answer", "placeholder", "duplicate_id", "unsafe_path", "wrong_slide",
                                    "unknown_field", "media_inventory", "coerced_type", "classification",
                                    "empty_lesson", "citation_hash", "duplicate_question"])
def test_malformed_approved_json_never_reaches_store(candidate, fresh_db, failure):
    pipeline, path, original = candidate
    payload = deepcopy(original)
    if failure == "truthy_approval":
        payload["approved"] = "true"
    elif failure == "legacy_schema":
        del payload["schema_version"]
    elif failure == "missing_review":
        del payload["approval"]
    elif failure == "missing_summary":
        del payload["concepts"][0]["summary"]
    elif failure == "wrong_answer":
        payload["question_pool"][0]["answer"] = "e"
    elif failure == "placeholder":
        payload["concepts"][0]["summary_status"] = "draft_placeholder"
    elif failure == "duplicate_id":
        payload["concepts"].append(deepcopy(payload["concepts"][0]))
    elif failure == "unsafe_path":
        payload["source_refs"][0]["source_path"] = "../RATIO.pptx"
    elif failure == "wrong_slide":
        payload["question_pool"][0]["source"]["slide_number"] = 99
    elif failure == "unknown_field":
        payload["ignored_problem"] = "must not be ignored"
    elif failure == "media_inventory":
        payload["media_ids"] = ["a" * 16]
    elif failure == "coerced_type":
        payload["source_refs"][0]["slide_count"] = "2"
    elif failure == "classification":
        payload["concepts"][0]["skill_id"] = "logic.pat.pattern_completion"
    elif failure == "empty_lesson":
        payload["concepts"][0]["learning_segments"] = []
        payload["concepts"][0]["examples"] = []
    elif failure == "citation_hash":
        payload["concepts"][0]["source_decks"][0]["source_hash"] = "a" * 64
    elif failure == "duplicate_question":
        payload["question_pool"].append(deepcopy(payload["question_pool"][0]))
    payload["content_sha256"] = content_digest(payload)
    path.write_text(json.dumps(payload), encoding="utf-8")
    assert main(["import", "--work", str(pipeline.work_dir), "--pack", str(path)]) == 2
    with db.read() as conn:
        assert conn.execute("SELECT count(*) FROM content_releases").fetchone()[0] == 0


def test_valid_import_and_retry_preserve_saved_payload(candidate, fresh_db, tmp_path, monkeypatch):
    pipeline, path, payload = candidate
    monkeypatch.setattr(settings, "media_dir", tmp_path / "installed-media")
    args = ["import", "--work", str(pipeline.work_dir), "--pack", str(path)]
    assert main(args) == 0
    assert main(args) == 0
    with db.read() as conn:
        rows = conn.execute("SELECT payload FROM content_releases").fetchall()
    assert len(rows) == 1 and json.loads(rows[0]["payload"]) == payload


def test_edited_content_cannot_bypass_review_by_rehashing(candidate):
    pipeline, _, payload = candidate
    payload["concepts"][0]["summary"] = "Edited outside the reviewed source draft."
    payload["content_sha256"] = content_digest(payload)
    validate_pack(payload)  # Structurally valid does not establish review provenance.
    with pytest.raises(ValueError, match="reviewed work directory"):
        prepare_import(payload, pipeline.work_dir)


def test_missing_or_malformed_reviewed_work_is_a_clear_refusal(candidate, tmp_path):
    pipeline, _, payload = candidate
    with pytest.raises(ValueError, match="requires the reviewed work directory"):
        prepare_import(payload, tmp_path / "missing-work")
    (pipeline.work_dir / "catalog.json").write_text("[]", encoding="utf-8")
    with pytest.raises(ValueError, match="invalid reviewed catalog"):
        prepare_import(payload, pipeline.work_dir)


def test_changed_source_and_changed_reviewed_catalog_block_import(candidate):
    pipeline, _, payload = candidate
    source = pipeline.source_dir / "RATIO.pptx"
    original = source.read_bytes()
    source.write_bytes(original + b"changed")
    with pytest.raises(ValueError, match="source missing, changed"):
        prepare_import(payload, pipeline.work_dir)
    source.write_bytes(original)
    path = pipeline.work_dir / "catalog.json"
    catalog = json.loads(path.read_text(encoding="utf-8"))
    catalog["decks"][0]["slides"][1]["notes"] = "Answer: B"
    path.write_text(json.dumps(catalog), encoding="utf-8")
    with pytest.raises(ValueError, match="reviewed work directory"):
        prepare_import(payload, pipeline.work_dir)


def test_candidate_retry_keeps_bytes_and_changed_content_needs_new_version(candidate):
    pipeline, path, payload = candidate
    original = path.read_bytes()
    assert pipeline.run_pack("1.0", True) == 0
    assert path.read_bytes() == original
    payload["concepts"][0]["summary"] = "Different lesson"
    payload["content_sha256"] = content_digest(payload)
    with pytest.raises(ValueError, match="immutable candidate conflict"):
        write_candidate(path, payload)
    assert path.read_bytes() == original


def jpeg(root, color="white"):
    root.mkdir(exist_ok=True)
    buffer = io.BytesIO()
    Image.new("RGB", (100, 50), color).save(buffer, format="JPEG")
    blob = buffer.getvalue()
    mid = hashlib.sha256(blob).hexdigest()[:16]
    (root / f"{mid}.jpg").write_bytes(blob)
    (root / "private-original.bin").write_bytes(b"private source")
    return mid, blob


def test_media_preflight_and_immutable_install(tmp_path):
    source, destination = tmp_path / "draft-media", tmp_path / "installed"
    mid, blob = jpeg(source)
    with pytest.raises(ValueError, match="missing"):
        install_media([mid, "a" * 16], source, destination)
    assert not destination.exists()
    install_media([mid], source, destination)
    install_media([mid], source, destination)
    assert [path.name for path in destination.iterdir()] == [f"{mid}.jpg"]
    assert (destination / f"{mid}.jpg").read_bytes() == blob
    (destination / f"{mid}.jpg").write_bytes(b"bad existing file")
    with pytest.raises(ValueError, match="hash/size mismatch"):
        install_media([mid], source, destination)
    assert (destination / f"{mid}.jpg").read_bytes() == b"bad existing file"


def test_correct_hash_of_non_jpeg_is_still_rejected(tmp_path):
    source = tmp_path / "draft-media"
    source.mkdir()
    blob = b"not a jpeg"
    mid = hashlib.sha256(blob).hexdigest()[:16]
    (source / f"{mid}.jpg").write_bytes(blob)
    with pytest.raises(ValueError, match="invalid delivery JPEG"):
        install_media([mid], source, tmp_path / "installed")
