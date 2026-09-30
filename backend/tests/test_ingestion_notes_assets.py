"""Recover solution images and hidden Office branches without merging roles."""
import copy
import io

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.util import Inches

from ingestion.pipeline import Pipeline


def image(slide, color):
    buffer = io.BytesIO()
    Image.new("RGB", (240, 100), color).save(buffer, format="PNG")
    buffer.seek(0)
    return slide.shapes.add_picture(buffer, Inches(1), Inches(4), Inches(3))


def add_notes_picture(slide):
    # NotesSlideShapes has no add_picture API. Use an image relationship plus
    # a native picture node, as actual corpus notes pages contain.
    staging = image(slide, "blue")
    node = copy.deepcopy(staging._element)
    part = slide.part.related_part(staging._element.blipFill.blip.rEmbed)
    note_rid = slide.notes_slide.part.relate_to(part, RT.IMAGE)
    node.blipFill.blip.rEmbed = note_rid
    node.nvPicPr.cNvPr.set("id", "301")
    slide.notes_slide.shapes._spTree.insert_element_before(node, "p:extLst")
    slide.shapes._spTree.remove(staging._element)


def extract(tmp_path, prs):
    source = tmp_path / "source"
    source.mkdir()
    path = source / "SPEED.pptx"
    prs.save(path)
    before = path.read_bytes()
    pipeline = Pipeline(source, tmp_path / "work")
    catalog = pipeline.normalize([path])
    report = pipeline.validate(catalog)
    drafts = pipeline.organize(catalog)
    assert path.read_bytes() == before
    assert not catalog["skipped"]
    return pipeline, catalog, report, drafts


def test_question_and_solution_picture_roles_survive_to_example(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(2))
    box.text_frame.text = "Find the distance travelled.\nA) 12 km\nB) 10 km"
    image(slide, "red")
    add_notes_picture(slide)
    slide.notes_slide.notes_text_frame.text = "Option A\nUse the solution diagram below."
    pipeline, catalog, report, drafts = extract(tmp_path, prs)
    result = catalog["decks"][0]["slides"][0]
    assert len(result["slide_media_ids"]) == len(result["notes_media_ids"]) == 1
    assert set(result["slide_media_ids"]).isdisjoint(result["notes_media_ids"])
    note_image = next(b for b in result["notes_blocks"] if b["type"] == "image")
    assert note_image["block_id"].startswith(result["slide_id"] + ":notes:")
    assert (pipeline.media_dir / note_image["original_file"]).is_file()
    assert note_image["surface"] == "notes"
    assert any(i["code"] == "visual_semantics_unresolved" and i.get("surface") == "notes" for i in result["issues"])
    assert report["ok"] and not report["ready_for_approval"]
    example = drafts["concepts"][0]["examples"][0]
    assert example["media_ids"] == result["slide_media_ids"]
    assert example["notes_media_ids"] == result["notes_media_ids"]
    assert any(b["block_id"] == note_image["block_id"] for b in example["notes_blocks"])
    assert {m["source"]["surface"] for m in drafts["concepts"][0]["media"]} == {"slide", "notes"}
    assert not any(i["code"] == "media_excluded" for i in drafts["review_issues"])


def test_notes_only_picture_is_kept_and_notes_page_furniture_is_explicit(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_notes_picture(slide)
    _, catalog, _, drafts = extract(tmp_path, prs)
    result = catalog["decks"][0]["slides"][0]
    assert result["kind"] == "explanation"
    assert result["notes_content_mode"] == "image_only"
    assert drafts["concepts"][0]["learning_segments"][0]["notes_media_ids"]
    furniture = [b for b in result["notes_blocks"] if b["type"] == "notes_furniture"]
    assert furniture and all(b["reason"] == "explicit_notes_page_placeholder" for b in furniture)
    assert not any(i["code"] == "unsupported_shape" for i in result["issues"])


def test_alternate_content_native_runs_and_xml_are_retained_without_guessing_branch(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(2))
    box.text_frame.text = "A hidden equation followed by minutes"
    mc = "http://schemas.openxmlformats.org/markup-compatibility/2006"
    alternate = etree.Element(f"{{{mc}}}AlternateContent", nsmap={"mc": mc, "a14": "http://schemas.microsoft.com/office/drawing/2010/main"})
    choice = etree.SubElement(alternate, f"{{{mc}}}Choice", Requires="a14")
    choice.append(copy.deepcopy(box._element))
    slide.shapes._spTree.remove(box._element)
    slide.shapes._spTree.append(alternate)
    _, catalog, report, _ = extract(tmp_path, prs)
    result = catalog["decks"][0]["slides"][0]
    branch = next(b for b in result["blocks"] if b["type"] == "alternate_content")
    assert "minutes" in " ".join(branch["native_runs"])
    assert "AlternateContent" in branch["xml"]
    assert branch["bounds_missing"]
    assert any(i["code"] == "alternate_content_unresolved" for i in result["issues"])
    assert report["ok"] and not report["ready_for_approval"]


def test_prior_catalogs_cannot_approve_without_notes_evidence(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(2))
    box.text_frame.text = "Distance is speed multiplied by time with matching measurement units."
    pipeline, catalog, _, _ = extract(tmp_path, prs)
    assert catalog["extraction_version"] == 3
    catalog["extraction_version"] = 2
    report = pipeline.validate(catalog)
    assert not report["ready_for_approval"]
    assert "catalog_reextract_required" in {i["code"] for i in report["review_items"]}


def test_static_embedded_object_preview_is_retained_without_reading_object_payload(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    staging = image(slide, "green")
    staging.crop_left = 0.1
    pic = copy.deepcopy(staging._element)
    part = slide.part.related_part(staging._element.blipFill.blip.rEmbed)
    rid = slide.notes_slide.part.relate_to(part, RT.IMAGE)
    pic.blipFill.blip.rEmbed = rid
    frame = parse_xml(f'''<p:graphicFrame xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
      xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
      xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
      <p:nvGraphicFramePr><p:cNvPr id="401" name="Static object preview"/><p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>
      <p:xfrm><a:off x="914400" y="3657600"/><a:ext cx="2743200" cy="1143000"/></p:xfrm>
      <a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/presentationml/2006/ole">
      <p:oleObj r:id="unused-object-id" progId="Word.Document.12"><p:embed/></p:oleObj>
      </a:graphicData></a:graphic></p:graphicFrame>''')
    frame.xpath(".//p:oleObj")[0].append(pic)
    slide.notes_slide.shapes._spTree.insert_element_before(frame, "p:extLst")
    slide.shapes._spTree.remove(staging._element)
    _, catalog, report, _ = extract(tmp_path, prs)
    result = catalog["decks"][0]["slides"][0]
    preview = next(b for b in result["notes_blocks"] if b.get("role") == "embedded_object_preview")
    assert preview["image_id"] in result["notes_media_ids"]
    assert preview["crop"]["left"] == 0.1
    assert any(i["code"] == "embedded_object_unresolved" for i in result["issues"])
    assert report["ok"] and not report["ready_for_approval"]
