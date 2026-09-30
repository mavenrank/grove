"""Ingestion must preserve solvable context and expose unresolved source loss."""
from __future__ import annotations

import io
import json
import shutil

import pytest
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.util import Inches

from ingestion.extract import inspect_image_blob
from ingestion.organize.questions import parse_question_slide
from ingestion.pipeline import Pipeline


@pytest.fixture
def pipeline(tmp_path):
    source = tmp_path / "source"
    source.mkdir()
    return Pipeline(source, tmp_path / "work")


def text(slide, value, x=1, y=1):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(7), Inches(1))
    shape.text_frame.text = value
    return shape


def picture(slide, size=(480, 80), label="distance = speed x time"):
    # A thin, highly compressed formula graphic was rejected by the old caps.
    image = Image.new("RGB", size, "white")
    ImageDraw.Draw(image).text((10, 15), label, fill="black")
    stream = io.BytesIO()
    image.save(stream, format="PNG")
    stream.seek(0)
    return slide.shapes.add_picture(stream, Inches(1), Inches(2), width=Inches(6))


def save(pipeline, prs, filename="SPEED.pptx"):
    path = pipeline.source_dir / filename
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return path


def run(pipeline):
    catalog = pipeline.normalize(pipeline.discover())
    report = pipeline.validate(catalog)
    return catalog, report, pipeline.organize(catalog)


def test_content_on_first_slide_and_empty_last_slide_survive(pipeline):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    text(slide, "1. Read the given distance.\n2. Divide by time.\nSpeed = Distance / Time")
    prs.slides.add_slide(prs.slide_layouts[6])
    save(pipeline, prs)
    catalog, report, organized = run(pipeline)
    deck = catalog["decks"][0]
    assert deck["slide_count"] == 2
    assert deck["slides"][0]["kind"] == "explanation"
    assert deck["slides"][1]["kind_reason"] == "empty"
    assert len(organized["decisions"]) == 2
    assert organized["decisions"][1]["outcome"] == "excluded"
    segment = organized["concepts"][0]["learning_segments"][0]
    assert "Speed = Distance / Time" in segment["raw_texts"][0]
    assert report["ok"]


def test_real_cover_is_measurable_not_position_alone(pipeline):
    prs = Presentation()
    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = "Speed"
    cover.placeholders[1].text = "A study introduction"
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    text(slide, "Distance is measured in kilometres and speed in kilometres per hour.")
    save(pipeline, prs)
    catalog, _, organized = run(pipeline)
    assert catalog["decks"][0]["slides"][0]["kind_reason"] == "short_title_placeholders"
    assert organized["decisions"][0]["reason"] == "short_title_placeholders"
    assert organized["decisions"][0]["outcome"] == "excluded"


def test_tables_grouped_steps_and_notes_keep_structure(pipeline):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    table_shape = slide.shapes.add_table(2, 3, Inches(1), Inches(1), Inches(6), Inches(2))
    values = [["Speed (km/h)", "Time (h)", "Distance (km)"], ["6", "2", "12"]]
    for row_index, row in enumerate(values):
        for col_index, value in enumerate(row):
            table_shape.table.cell(row_index, col_index).text = value
    group = slide.shapes.add_group_shape()
    step = group.shapes.add_textbox(Inches(1), Inches(4), Inches(5), Inches(1))
    step.text_frame.text = "1. Match units\n2. Multiply speed by time"
    step.text_frame.paragraphs[1].level = 1
    step.text_frame.paragraphs[1].runs[0].font.bold = True
    group.left, group.top, group.width = Inches(2), Inches(3), Inches(6)
    slide.notes_slide.notes_text_frame.text = "Check units first.\nThen substitute the givens."
    save(pipeline, prs)
    catalog, report, _ = run(pipeline)
    result = catalog["decks"][0]["slides"][0]
    table = next(b for b in result["blocks"] if b["type"] == "table")
    assert [[c["text"] for c in row] for row in table["rows"]] == values
    block = next(b for b in result["blocks"] if b["shape_id"] == step.shape_id)
    assert block["group_path"] == [group.shape_id]
    assert block["bounds_emu"]["left"] == Inches(2)
    assert block["paragraphs"][1]["level"] == 1
    assert block["paragraphs"][1]["runs"][0]["bold"] is True
    assert "\n" in result["notes"]
    assert block["block_id"].startswith(result["slide_id"])
    assert report["ok"]


def test_native_math_and_superscript_are_never_silently_flattened(pipeline):
    from lxml import etree

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = text(slide, "Distance = Speed x Time")
    shape.text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr().set("baseline", "30000")
    shape._element.append(etree.fromstring(
        b'<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>x=2</m:t></m:r></m:oMath>'
    ))
    save(pipeline, prs)
    catalog, report, organized = run(pipeline)
    result = catalog["decks"][0]["slides"][0]
    assert result["blocks"][0]["paragraphs"][0]["runs"][0]["baseline"] == "30000"
    assert "native_equation_unresolved" in {i["code"] for i in result["issues"]}
    with pytest.raises(ValueError, match="native_equation_unresolved"):
        pipeline.pack(catalog, report, organized, "test", True)


def test_tiny_formula_is_retained_and_original_saved(pipeline):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    pic = picture(slide)
    pic.crop_left = 0.1
    save(pipeline, prs)
    catalog, report, organized = run(pipeline)
    result = catalog["decks"][0]["slides"][0]
    assert result["content_mode"] == "image_only"
    block = result["blocks"][0]
    assert block["original_size"] == [480, 80]
    assert block["crop"]["left"] == 0.1
    assert (pipeline.media_dir / block["original_file"]).is_file()
    assert (pipeline.media_dir / f"{block['image_id']}.jpg").is_file()
    assert organized["concepts"][0]["media"][0]["image_id"] == block["image_id"]
    assert report["ok"] and not report["ready_for_approval"]
    with pytest.raises(ValueError, match="visual_semantics_unresolved"):
        pipeline.pack(catalog, report, organized, "test", True)


def test_image_led_first_slide_is_not_a_cover(pipeline):
    prs = Presentation()
    picture(prs.slides.add_slide(prs.slide_layouts[6]))
    prs.slides.add_slide(prs.slide_layouts[6])
    save(pipeline, prs)
    catalog, _, organized = run(pipeline)
    assert catalog["decks"][0]["slides"][0]["kind"] == "explanation"
    assert organized["concepts"][0]["learning_segments"][0]["media_ids"]


def test_repeated_diagram_survives_three_distinct_sources(pipeline):
    for n in range(3):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        text(slide, f"Use this relation to find an unknown value in example {n}.")
        picture(slide)
        save(pipeline, prs, f"{n}-SPEED.pptx")
    _, _, organized = run(pipeline)
    assert len(organized["repeated_media"]) == 1
    assert organized["repeated_media"][0]["distinct_sources"] == 3
    segments = organized["concepts"][0]["learning_segments"]
    assert len(segments) == 3 and all(s["media_ids"] for s in segments)


def test_hybrid_question_keeps_diagram_with_answer_and_snapshot(pipeline):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    text(slide, "Question 1\nUsing the diagram, find the distance in kilometres.\nA) 12\nB) 6")
    picture(slide)
    slide.notes_slide.notes_text_frame.text = "Answer: A. Multiply speed by time."
    save(pipeline, prs)
    catalog, report, organized = run(pipeline)
    example = organized["concepts"][0]["examples"][0]
    assert example["answer"] == "a"
    assert example["answer_status"] == "notes_confirmed"
    assert example["media_ids"]
    assert any(b["type"] == "image" and b["image_id"] == example["media_ids"][0] for b in example["blocks"])
    assert example["source"]["source_hash"] == catalog["decks"][0]["source_hash"]
    assert example["source"]["slide_id"]
    assert not report["ready_for_approval"]


def test_same_stem_in_two_folders_cannot_overwrite_media(pipeline):
    for folder, label in [("a", "Speed = distance / time"), ("b", "Distance = speed x time")]:
        prs = Presentation()
        picture(prs.slides.add_slide(prs.slide_layouts[6]), label=label)
        save(pipeline, prs, f"{folder}/SPEED.pptx")
    catalog, _, _ = run(pipeline)
    ids = [d["deck_id"] for d in catalog["decks"]]
    assert len(set(ids)) == 2
    assert set(catalog["media_index"]) == set(ids)
    assert catalog["media_index"][ids[0]] != catalog["media_index"][ids[1]]


def test_path_identity_is_stable_but_slide_snapshot_changes(pipeline):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    text(slide, "Distance equals speed multiplied by elapsed time.")
    save(pipeline, prs)
    first = run(pipeline)[0]["decks"][0]
    slide.shapes[0].text_frame.text = "Speed equals distance divided by elapsed time."
    save(pipeline, prs)
    second = run(pipeline)[0]["decks"][0]
    assert first["deck_id"] == second["deck_id"]
    assert first["source_hash"] != second["source_hash"]
    assert first["slides"][0]["slide_id"] != second["slides"][0]["slide_id"]


def test_unparsed_question_and_vector_shape_are_reviewable(pipeline):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    text(slide, "Question 1\nFind the missing value from this diagram")
    slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1), Inches(2), Inches(2), Inches(3))
    save(pipeline, prs)
    catalog, report, organized = run(pipeline)
    codes = {i["code"] for i in organized["review_issues"]}
    assert {"unsupported_shape", "question_parse_failed"} <= codes
    assert organized["decisions"][0]["outcome"] == "needs_review"
    with pytest.raises(ValueError, match="unresolved review"):
        pipeline.pack(catalog, report, {"concepts": [], "review_issues": [], "stats": {}}, "test", True)


@pytest.mark.parametrize("options,answer,code", [
    ("A) 12\nB) 6", "E", "answer_outside_options"),
    ("B) 12\nC) 6", "B", "option_mapping_ambiguous"),
    ("A)\nB) 6\n12", "A", "option_mapping_ambiguous"),
])
def test_question_answer_cannot_confirm_an_invalid_option_mapping(options, answer, code):
    question = parse_question_slide({"texts": [f"Question 1\nFind distance for a speed of six and two hours.\n{options}"],
                                     "notes": f"Answer: {answer}. Multiply the values.",
                                     "deck_id": "sample", "slide_number": 1, "source_file": "SPEED.pptx"})
    assert question["answer"] is None
    assert code in {i["code"] for i in question["issues"]}


def test_uppercase_extension_hidden_source_and_placeholder_are_accounted(pipeline):
    prs = Presentation()
    text(prs.slides.add_slide(prs.slide_layouts[6]), "Speed is the distance travelled per unit of elapsed time.")
    source = save(pipeline, prs, "SPEED.PPTX")
    shutil.copyfile(source, pipeline.source_dir / ".hidden.pptx")
    (pipeline.source_dir / "SPEED.pdf").write_bytes(b"placeholder")
    catalog, report, _ = run(pipeline)
    assert catalog["file_count"] == 3
    assert len(catalog["decks"]) == 2
    assert catalog["skipped"][0]["reason"] == "hidden_source"
    assert "format_not_extracted" in {i["code"] for i in report["review_items"]}
    assert not report["ready_for_approval"]


def test_corrupt_and_oversized_images_have_explicit_reasons():
    assert inspect_image_blob(b"bad image")[1][0]["code"] == "image_decode_failed"
    assert inspect_image_blob(b"x" * (8 * 1024 * 1024 + 1))[1][0]["code"] == "image_byte_limit"


def test_missing_delivery_media_is_a_hard_failure(pipeline):
    prs = Presentation()
    picture(prs.slides.add_slide(prs.slide_layouts[6]))
    save(pipeline, prs)
    catalog, _, _ = run(pipeline)
    mid = catalog["decks"][0]["slides"][0]["media_ids"][0]
    (pipeline.media_dir / f"{mid}.jpg").unlink()
    report = pipeline.validate(catalog)
    assert not report["ok"] and "missing/invalid media" in report["problems"][0]


@pytest.mark.parametrize("failure", ["corrupt", "mismatched_index"])
def test_media_integrity_is_checked_again_at_approval(pipeline, failure):
    prs = Presentation()
    picture(prs.slides.add_slide(prs.slide_layouts[6]))
    save(pipeline, prs)
    catalog, report, organized = run(pipeline)
    deck = catalog["decks"][0]
    if failure == "corrupt":
        mid = deck["slides"][0]["media_ids"][0]
        (pipeline.media_dir / f"{mid}.jpg").write_bytes(b"corrupted media")
    else:
        catalog["media_index"][deck["deck_id"]][1] = []
    assert not pipeline.validate(catalog)["ok"]
    with pytest.raises(ValueError, match="validation failures"):
        pipeline.pack(catalog, report, organized, "test", True)


def test_legacy_catalog_cannot_approve_without_reextracting(pipeline):
    prs = Presentation()
    text(prs.slides.add_slide(prs.slide_layouts[6]), "Speed is the distance travelled per unit of elapsed time.")
    save(pipeline, prs)
    catalog, report, organized = run(pipeline)
    del catalog["extraction_version"]
    with pytest.raises(ValueError, match="catalog_reextract_required"):
        pipeline.pack(catalog, report, organized, "test", True)


def test_generic_summary_cannot_be_approved_as_a_ready_lesson(pipeline):
    prs = Presentation()
    text(prs.slides.add_slide(prs.slide_layouts[6]), "Speed")
    save(pipeline, prs)
    catalog, report, organized = run(pipeline)
    assert report["ok"]
    assert organized["concepts"][0]["summary_status"] == "draft_placeholder"
    with pytest.raises(ValueError, match="summary_needs_review"):
        pipeline.pack(catalog, report, {"concepts": [{"summary": "Fake ready lesson"}], "stats": {}}, "test", True)


def test_validation_recomputes_duplicates_and_detects_id_collision(pipeline):
    prs = Presentation()
    text(prs.slides.add_slide(prs.slide_layouts[6]), "Speed is the distance travelled per unit of elapsed time.")
    first = save(pipeline, prs)
    shutil.copyfile(first, pipeline.source_dir / "SPEED-copy.pptx")
    catalog, _, _ = run(pipeline)
    assert pipeline.validate(catalog)["counts"]["duplicates"] == 1
    catalog["decks"][1]["source_hash"] = "changed"
    assert pipeline.validate(catalog)["counts"]["duplicates"] == 0
    catalog["decks"][1]["deck_id"] = catalog["decks"][0]["deck_id"]
    assert "duplicate deck identity" in pipeline.validate(catalog)["problems"][0]


def test_example_limit_keeps_full_review_candidates(pipeline):
    prs = Presentation()
    for index in range(23):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        text(slide, f"Question {index + 1}\nFind distance for speed six and duration two in example {index}.\nA) 12\nB) 6")
        slide.notes_slide.notes_text_frame.text = "Answer: A. Multiply speed by duration."
    save(pipeline, prs)
    _, _, organized = run(pipeline)
    concept = organized["concepts"][0]
    assert len(concept["examples"]) == 20
    assert len(concept["review_candidates"]["examples"]) == 23
    assert {"field": "examples", "retained": 20, "candidates": 23, "reason": "presentation_limit"} in concept["truncations"]


def test_review_report_accounts_for_source_slides(pipeline):
    prs = Presentation()
    picture(prs.slides.add_slide(prs.slide_layouts[6]))
    prs.slides.add_slide(prs.slide_layouts[6])
    save(pipeline, prs)
    assert pipeline.run_all() == 0  # draft extraction succeeds, approval remains blocked
    report = json.loads((pipeline.work_dir / "review-report.json").read_text())
    assert len(report["slide_decisions"]) == 2
    assert report["limits"]["preview_dimension"] == 1400
    assert not report["ready_for_approval"]
    assert report["validation_review"][0]["code"] == "visual_semantics_unresolved"
