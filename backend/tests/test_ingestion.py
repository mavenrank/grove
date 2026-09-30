"""Ingestion pipeline tests: classification, extraction, validation, approval gate."""
from __future__ import annotations

import shutil
import tempfile
from pathlib import Path

import pytest

from ingestion.config import classify_title
from ingestion.pipeline import Pipeline, slugify


@pytest.fixture()
def workroot():
    d = Path(tempfile.mkdtemp(prefix="grove-ingest-"))
    yield d
    shutil.rmtree(d, ignore_errors=True)


def test_classify_titles():
    cases = {
        "31-PERCENTAGES_1-2023-04-28": "quant.frp.percentages",
        "21-BLOOD_RELATIONS_1-2023-04-19": "logic.rel.blood_relations",
        "47-SIMPLE_INTEREST-2023-05-20": "quant.avg.simple_interest",
        "59-NUMBERS_4_HCF_AND_LCM-2023-05-31": "quant.numbers.hcf_lcm",
        "015-DATA_ARRANGEMENTS_1-2023-04-10": "logic.arr.data_arrangements",
        "35-VOCABULARY_SYNONYMS-2023-05-03": ("verbal.vocab.synonyms", "deferred"),
        "16-Karatsuba Algorithm-19-08-2024": ("other.algorithms", "deferred"),
    }
    for title, expected in cases.items():
        skill, status = classify_title(title)
        if isinstance(expected, tuple):
            assert (skill, status) == expected, title
        else:
            assert skill == expected and status == "active", title


def _make_deck(path: Path, title_text: str, body_text: str) -> None:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title_text
    slide.placeholders[1].text = body_text
    prs.save(path)


def test_extract_pptx_and_dedupe(workroot):
    src = workroot / "src"
    src.mkdir()
    _make_deck(src / "01-PERCENTAGES_1.pptx", "Percentages", "x% of N = (x/100) * N")
    _make_deck(src / "02-PERCENTAGES_1-copy.pptx", "Percentages", "x% of N = (x/100) * N")

    p = Pipeline(src, workroot / "work")
    files = p.discover()
    assert len(files) == 2
    catalog = p.normalize(files)
    deck = catalog["decks"][0]
    assert deck["slides"][0]["title"] == "Percentages"
    assert deck["slides"][0]["texts"]
    assert deck["source_hash"]

    report = p.validate(catalog)
    assert report["counts"]["duplicates"] == 1
    assert report["counts"]["unique"] == 1
    assert report["ok"] is True

    concepts = p.organize(catalog)["concepts"]
    assert len(concepts) == 1
    assert concepts[0]["skill_id"] == "quant.frp.percentages"
    assert concepts[0]["source_decks"][0]["source_hash"]


def test_pack_requires_approval(workroot):
    src = workroot / "src"
    src.mkdir()
    _make_deck(src / "01-RATIO.pptx", "Ratios",
               "A ratio a:b divides a whole into a plus b equal parts; one part equals total divided by the sum.")

    p = Pipeline(src, workroot / "work")
    catalog = p.normalize(p.discover())
    report = p.validate(catalog)
    organized = p.organize(catalog)

    draft = p.pack(catalog, report, organized, "0.1.0", approved=False)
    assert draft["concepts"] == []
    assert draft["approved"] is False

    approved = p.pack(catalog, report, organized, "0.1.0", approved=True)
    assert len(approved["concepts"]) == 1
    assert approved["source_refs"][0]["source_hash"]


def test_slugify():
    assert slugify("12-Percent Change! (2)") == "12-percent-change-2"
