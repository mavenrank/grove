"""Bench regressions test losses between stages, not implementation mirrors."""
from __future__ import annotations

import copy
import json
import sqlite3

import pytest
from pptx import Presentation
from pptx.util import Inches

from ingestion.bench import analyze, pptx_probe, read_release, run_bench
from ingestion.bench_report import render_report


def fixture_data():
    src = {"deck_id": "speed", "source_path": "SPEED.pptx", "source_hash": "abc", "slide_count": 1}
    concept = {"id": "concept.quant.tws.speed_distance_time", "title": "Speed", "skill_id": "quant.tws.speed_distance_time",
               "summary": "Teaching material assembled from 0 explanation slides", "formulas": ["Speed = ?"],
               "examples": [], "media": [], "source_decks": [src]}
    slide = {"slide_number": 1, "kind": "question", "texts": ["Find distance", "A) 12", "B) 10"],
             "issues": [], "content_mode": "text_only"}
    deck = {**src, "status": "active", "skill_id": concept["skill_id"], "slides": [slide]}
    example = {"prompt": "Find distance", "options": {"a": "12", "b": "10"}, "answer": "a",
               "explanation": "Multiply speed and time", "source": src, "blocks": [{"type": "text"}],
               "media_ids": ["1234567890abcdef"], "answer_status": "notes_confirmed"}
    draft = {**concept, "summary_status": "draft_placeholder", "examples": [example],
             "learning_segments": [{"texts": ["Distance = speed × time"], "source": src}]}
    return ({"release_id": "test", "version": "1", "concepts": [concept]},
            {"decks": [deck], "skipped": []},
            {"concepts": [draft], "question_pool": [], "decisions": [], "review_issues": []},
            {"ok": True, "ready_for_approval": False, "problems": []})


def test_all_topics_have_stage_specific_loss_evidence_and_stable_ids():
    payload, catalog, drafts, validation = fixture_data()
    payload["concepts"].append({"id": "authored", "title": "Authored", "skill_id": "none", "summary": "Review", "source_decks": []})
    report = analyze(payload, catalog, drafts, validation, {})
    assert len(report["topics"]) == 2
    codes = report["counts"]["codes"]
    assert codes["source_refs_absent"] == 1
    assert codes["formula_question_fronts"] == 1
    assert codes["draft_recovers_examples"] == 1
    assert codes["public_field_dropped"] == 1
    assert codes["public_example_field_dropped"] == 3
    assert report["topics"][0]["independently_verified_by_bench"] == 0
    repeated = analyze(payload, catalog, drafts, validation, {})
    assert [f["id"] for f in report["findings"]] == [f["id"] for f in repeated["findings"]]


def test_native_probe_detects_missing_text_and_slide_count_with_citations():
    payload, catalog, drafts, validation = fixture_data()
    original = {"speed.pptx": [{"slide_number": 1, "native_runs": ["Distance = speed × time"]},
                              {"slide_number": 2, "native_runs": ["A missing slide"]}]}
    report = analyze(payload, catalog, drafts, validation, original)
    assert report["counts"]["codes"]["source_slide_loss"] == 1
    losses = [f for f in report["findings"] if f["code"] == "native_text_loss"]
    assert {f["source"]["slide_number"] for f in losses} == {1, 2}
    assert all(f["source"]["source_hash"] == "abc" for f in losses)


def test_read_release_selects_app_latest_without_changing_database(tmp_path):
    path = tmp_path / "store.db"
    with sqlite3.connect(path) as conn:
        conn.execute("CREATE TABLE content_releases(release_id TEXT,version TEXT,created_at TEXT,payload TEXT)")
        for version, time in (("9", "2020"), ("1", "2021")):
            conn.execute("INSERT INTO content_releases VALUES(?,?,?,?)", ("test", version, time, json.dumps({"version": version})))
    before = path.read_bytes()
    payload, metadata = read_release(path)
    assert payload["version"] == metadata["version"] == "1"
    assert path.read_bytes() == before
    with pytest.raises(sqlite3.OperationalError):
        read_release(tmp_path / "missing.db")
    assert not (tmp_path / "missing.db").exists()


def test_fresh_bench_reextracts_real_package_without_modifying_original(tmp_path):
    source = tmp_path / "source"
    source.mkdir()
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(3))
    box.text_frame.text = "Distance is speed multiplied by time. Keep the units consistent."
    path = source / "SPEED.pptx"
    prs.save(path)
    before = path.read_bytes()
    payload, _, _, _ = fixture_data()
    payload["concepts"][0]["source_decks"][0]["source_hash"] = __import__("hashlib").sha256(before).hexdigest()
    probe = pptx_probe(path)
    assert len(probe) == 1 and "Keep the units consistent." in probe[0]["native_runs"][0]
    work = tmp_path / "bench"
    report = run_bench(payload, source, work)
    assert report["counts"]["topics"] == 1
    assert report["topics"][0]["sources_extracted"] == 1
    assert report["topics"][0]["draft"]["learning_segments"] == 1
    assert "native_text_loss" not in report["counts"]["codes"]
    assert path.read_bytes() == before
    assert (work / "bench-report.html").exists()
    with pytest.raises(ValueError, match="new or empty"):
        run_bench(payload, source, work)
    with pytest.raises(ValueError, match="outside"):
        run_bench(payload, source, source / "work")


def test_outside_and_missing_source_refs_are_reported_not_read(tmp_path):
    source = tmp_path / "source"
    source.mkdir()
    payload, _, _, _ = fixture_data()
    payload["concepts"][0]["source_decks"][0]["source_path"] = "../outside.pptx"
    report = run_bench(payload, source, tmp_path / "bench")
    assert report["topics"][0]["sources_extracted"] == 0
    assert report["counts"]["codes"]["source_unavailable"] == 1
    assert report["skipped_sources"][0]["reason"] == "missing_or_outside_source_root"


def test_heuristics_do_not_claim_bad_math_and_report_escapes_source_markup():
    payload, catalog, drafts, validation = fixture_data()
    payload["concepts"][0]["summary"] = "<script>alert(1)</script>(" + "a" * 510
    before = copy.deepcopy(payload)
    report = analyze(payload, catalog, drafts, validation, {})
    formatting = [f for f in report["findings"] if f["stage"] == "formatting"]
    assert {f["code"] for f in formatting} == {"delimiter_imbalance", "long_text_blob"}
    assert all(f["confidence"] == "heuristic" for f in formatting)
    rendered = render_report(report)
    assert "<script>" not in rendered and "&lt;script&gt;" in rendered
    assert payload == before
