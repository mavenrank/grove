"""Source-format adapter registry (handoff §13 extension point).

The pipeline discovers files and delegates extraction to the adapter
registered for the file's extension. Adding a format (e.g. real PDF/DOCX
text extraction, or an OCR stage over slide images) is a registration, not a
pipeline edit:

    from .extractors import register_extractor

    @register_extractor(".pdf")
    def extract_pdf(path: Path, ctx: ExtractionContext):
        return slides, media_by_slide

Adapter contract — returns `(slides, media_by_slide)`:
  slides          list of {slide_number, title, texts, notes, kind, media_ids}
  media_by_slide  {slide_number: [image_id, ...]} referencing work-dir media

Adapters receive untrusted input: enforce size bounds, never execute embedded
content, and never raise for a single bad file (the pipeline records skips).
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable

Extractor = Callable[[Path, "ExtractionContext"], tuple[list[dict[str, Any]], dict[int, list[str]]]]

_EXTRACTORS: dict[str, Extractor] = {}

SUPPORTED_SUFFIXES = (".pptx", ".pdf", ".docx")


@dataclass
class ExtractionContext:
    """Everything an adapter may need from the pipeline run."""

    media_dir: Path            # where normalized images are written
    source_dir: Path           # the corpus root (for relative provenance)

    def save_media(self, processed: dict[str, Any] | None) -> str | None:
        """Persist one processed image blob; content-hash dedupe."""
        if not processed:
            return None
        image_id = processed["image_id"]
        media_path = self.media_dir / f"{image_id}.jpg"
        if not media_path.exists():
            media_path.write_bytes(processed["data"])
        return image_id


def register_extractor(suffix: str):
    """Register an adapter for a file extension (e.g. '.pptx')."""

    def deco(fn: Extractor) -> Extractor:
        _EXTRACTORS[suffix.lower()] = fn
        return fn

    return deco


def get_extractor(suffix: str) -> Extractor | None:
    return _EXTRACTORS.get(suffix.lower())


def supported(suffix: str) -> bool:
    return suffix.lower() in _EXTRACTORS


# ---------------------------------------------------------------------------
# Built-in adapters
# ---------------------------------------------------------------------------

def extract_pptx(path: Path, ctx: ExtractionContext) -> tuple[list[dict[str, Any]], dict[int, list[str]]]:
    """Deep PPTX extraction: recursed text, notes, slide kind, and media files.

    Pictures (the corpus renders its teaching content as slide images) are
    normalized into work/media/<image_id>.jpg and referenced by id.
    """
    from pptx import Presentation

    from .extract import extract_slide_content, process_image_blob

    prs = Presentation(str(path))
    total = len(prs.slides)
    slides: list[dict[str, Any]] = []
    media_by_slide: dict[int, list[str]] = {}
    for idx, slide in enumerate(prs.slides, start=1):
        # slide 1 is a provider cover only in multi-slide decks; a lone
        # slide is content
        first_of_many = idx == 1 and total > 1
        content = extract_slide_content(slide, first_of_many=first_of_many)
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = _clean(slide.notes_slide.notes_text_frame.text)
        media_ids: list[str] = []
        for img in content["image_blobs"]:
            mid = ctx.save_media(process_image_blob(img["blob"]))
            if mid:
                media_ids.append(mid)
        if media_ids:
            media_by_slide[idx] = media_ids
        title = ""
        try:
            if slide.shapes.title is not None:
                title = _clean(slide.shapes.title.text)
        except Exception:
            pass
        if content["texts"] or notes or media_ids:
            slides.append({
                "slide_number": idx,
                "title": title,
                "texts": content["texts"],
                "notes": notes,
                "kind": content["kind"],
                "media_ids": media_ids,
            })
    return slides, media_by_slide


register_extractor(".pptx")(extract_pptx)


def extract_generic(path: Path, ctx: ExtractionContext) -> tuple[list[dict[str, Any]], dict[int, list[str]]]:
    """PDF/DOCX placeholder: preserve as single-record sources (deep extraction later)."""
    return ([{
        "slide_number": 0,
        "title": path.stem,
        "texts": [f"[{path.suffix.lstrip('.').upper()} source, not text-extracted in v1]"],
        "notes": "",
        "kind": "other",
        "media_ids": [],
    }], {})


for _sfx in (".pdf", ".docx"):
    register_extractor(_sfx)(extract_generic)


_WS = re.compile(r"\s+")


def _clean(text: str | None) -> str:
    if not text:
        return ""
    return _WS.sub(" ", text).strip()
